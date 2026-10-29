"""Shared test fixtures (no AWS):

  settings         Settings with a tmp data_dir (no .env)
  fixtures_dir     Path to tests/fixtures
  fake_llm         FakeLLM (empty queue; set .tool_responses)
"""
from __future__ import annotations

from pathlib import Path

import pytest

from adr_generator.config import Settings
from adr_generator.llm.bedrock import FakeLLM

FIXTURES = Path(__file__).resolve().parent / "fixtures"


def make_settings(tmp_path: Path, **overrides) -> Settings:
    base = dict(DATA_DIR=str(tmp_path / "data"))
    base.update(overrides)
    return Settings(_env_file=None, **base)


@pytest.fixture
def settings(tmp_path) -> Settings:
    return make_settings(tmp_path)


@pytest.fixture
def fixtures_dir() -> Path:
    return FIXTURES


@pytest.fixture
def fake_llm() -> FakeLLM:
    return FakeLLM()


def make_docx(path: Path) -> None:
    import docx

    d = docx.Document()
    d.add_heading("Orders architecture", level=1)
    d.add_paragraph("Customers reach an ALB which forwards to ECS.")
    t = d.add_table(rows=2, cols=3)
    for i, h in enumerate(["Component", "Service", "Notes"]):
        t.rows[0].cells[i].text = h
    for i, v in enumerate(["db", "Amazon RDS", "multi-AZ"]):
        t.rows[1].cells[i].text = v
    d.save(str(path))
