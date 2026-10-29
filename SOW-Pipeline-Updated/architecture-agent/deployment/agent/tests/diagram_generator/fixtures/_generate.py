"""Regenerates the binary fixtures: b_prose.pdf, c_tables.docx, e_scanned.pdf.

a_clean.md, d_vague.md, expectations.json and README.md are hand-written text files.

Run from this folder:
  uv run --python 3.12 --with reportlab --with python-docx --with pillow python _generate.py

All companies, people and systems are fictional.
"""
from __future__ import annotations

import os
import random
from pathlib import Path

HERE = Path(__file__).resolve().parent


# --------------------------------------------------------------------------
# b_prose.pdf  - architecture hidden in prose, text-based PDF
# --------------------------------------------------------------------------
def make_b_prose() -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
    )

    styles = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=styles["BodyText"], fontName="Helvetica",
                          fontSize=10.5, leading=15, spaceAfter=7)
    h1 = ParagraphStyle("h1", parent=styles["Heading1"], fontName="Helvetica-Bold",
                        fontSize=15, spaceBefore=12, spaceAfter=8)
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontName="Helvetica-Bold",
                        fontSize=12, spaceBefore=8, spaceAfter=5)
    title = ParagraphStyle("title", parent=styles["Title"], fontName="Helvetica-Bold",
                           fontSize=22, spaceAfter=10)
    sub = ParagraphStyle("sub", parent=body, fontSize=11, textColor=colors.HexColor("#444444"))

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.drawString(20 * mm, 12 * mm,
                          "Brightpath Cloud Consulting (fictional) - Solution Design Document v1.0 FINAL - Confidential")
        canvas.drawRightString(190 * mm, 12 * mm, "Page %d" % doc.page)
        canvas.restoreState()

    S = []
    P = lambda t, st=body: S.append(Paragraph(t, st))

    P("Solution Design Document", title)
    P("Order Processing Modernisation Programme", sub)
    P("Prepared for: Lumen Home Goods (the customer) &nbsp;|&nbsp; Prepared by: Brightpath Cloud Consulting "
      "&nbsp;|&nbsp; Version 1.0 FINAL", sub)
    S.append(Spacer(1, 6))

    P("1. Executive Summary", h1)
    P("Lumen Home Goods is a mid-sized online retailer of furniture and home accessories. Over the last two years "
      "the customer has seen order volumes grow by roughly 140 percent, and the existing monolithic order "
      "system, hosted on a single rented server, can no longer absorb seasonal peaks without manual intervention. "
      "This document describes the finalised solution that Brightpath Cloud Consulting will deliver to replace "
      "that system with a serverless, event-driven order processing platform on Amazon Web Services.")
    P("The programme is sponsored by the customer's Head of Digital and is expected to reduce checkout failures "
      "during peak periods, shorten the time from order placement to confirmation email, and remove the "
      "operational burden of patching and capacity planning. The design has been reviewed with the customer's "
      "engineering lead and the finance team and is considered final for the purposes of the delivery "
      "plan and commercial terms in this document.")

    P("2. Business Context and Objectives", h1)
    P("Black Friday and the pre-holiday period account for nearly 45 percent of the customer's annual revenue. "
      "During the last peak, roughly one in twenty checkout attempts timed out. The customer has set four "
      "objectives for the programme: sustain ten times the current peak order rate without degradation, "
      "confirm at least 99 percent of valid orders within 30 seconds, keep the cost per order flat as volume "
      "grows, and give the customer's support staff a reliable record of the state of every order.")
    P("Success will be measured by the checkout completion rate, the median and 99th percentile "
      "order-confirmation latency, and the number of manual interventions per thousand orders. Baselines "
      "were captured during discovery workshops held in March and April.")

    P("3. Scope", h1)
    P("In scope are the customer-facing storefront web application, customer sign-up and sign-in, the order "
      "intake API, asynchronous order processing including inventory reservation and payment, and "
      "transactional email notifications. Out of scope are native mobile applications, replacement of the "
      "customer's ERP and warehouse management systems, migration of historical orders older than 24 "
      "months, and any change to the product catalogue management process. The customer will remain "
      "responsible for its agreement with the external payment provider.")

    P("4. Proposed Solution Architecture", h1)
    P("The platform is deployed in a single AWS Region, Europe (Ireland) eu-west-1. It follows a "
      "serverless pattern throughout: there are no servers, containers or virtual networks for the delivery "
      "team to operate. The remainder of this section walks through the solution from the shopper's browser "
      "to the back-end workflow.")

    P("4.1 Storefront delivery", h2)
    P("Shoppers reach the storefront through their web browser. The storefront is a static single-page "
      "application whose build artefacts are stored in a private Amazon S3 bucket. Amazon CloudFront is "
      "placed in front of the bucket and serves the application to shoppers worldwide; CloudFront reads from "
      "the bucket through an origin access control so that the bucket itself is never publicly readable. "
      "The same CloudFront distribution also fronts the API so that the browser talks to a single domain.")

    P("4.2 Identity", h2)
    P("Customer accounts are managed by Amazon Cognito. A Cognito user pool handles registration, sign-in, "
      "password recovery and multi-factor authentication. After signing in, the browser application receives "
      "a JSON Web Token from Amazon Cognito and presents it on every API call.")

    P("4.3 Order intake", h2)
    P("The browser calls a REST API exposed by Amazon API Gateway. API Gateway validates the token with a "
      "Cognito authorizer and invokes an AWS Lambda function named the Order API function. This function "
      "validates the basket, writes a new order record with the status PENDING to an Amazon DynamoDB table "
      "called Orders, and then places an order message on an Amazon SQS queue. A dead-letter queue "
      "captures messages that fail processing after three attempts. The function returns an order "
      "identifier to the browser immediately, so the shopper never waits on payment processing.")

    P("4.4 Asynchronous processing and orchestration", h2)
    P("A second AWS Lambda function, the Order Processor, consumes messages from the Amazon SQS queue and "
      "starts one AWS Step Functions Standard workflow execution per order. The workflow performs the "
      "business steps in a fixed order. First it reserves stock by updating an Inventory table in Amazon "
      "DynamoDB using the native Step Functions service integration. Second, it invokes a Payment Lambda "
      "function that calls the HTTPS API of PayFlow, the external payment provider, to authorise and capture "
      "the payment. Third, it updates the order record in the Orders table to CONFIRMED. Finally, it publishes "
      "an OrderConfirmed event to a custom Amazon EventBridge event bus. If the payment is declined, the "
      "workflow releases the reserved stock and sets the order status to FAILED.")

    P("4.5 Notifications and events", h2)
    P("An Amazon EventBridge rule matches OrderConfirmed events and invokes a Notification Lambda function. "
      "That function renders the confirmation email and sends it to the shopper through Amazon SES. Using "
      "Amazon EventBridge keeps the workflow decoupled from notifications, so that additional consumers, such "
      "as a future loyalty-points service, can subscribe later without any change to the order workflow.")

    P("4.6 Observability and security", h2)
    P("All Lambda functions and the workflow emit logs and metrics to Amazon CloudWatch, and the customer's "
      "support staff will use a CloudWatch dashboard and alarms on queue depth and workflow failures. Each "
      "Lambda function runs with its own least-privilege IAM role, and the PayFlow API credentials are "
      "held in AWS Secrets Manager and read by the Payment Lambda function at runtime.")

    P("5. Delivery Timeline", h1)
    data = [
        ["Phase", "Duration", "Key outputs"],
        ["Mobilisation and design sign-off", "2 weeks", "Final solution design, delivery plan"],
        ["Platform foundation and storefront", "4 weeks", "Static site, identity, CI/CD pipelines"],
        ["Order intake and processing", "6 weeks", "API, queue, workflow, payment integration"],
        ["Notifications, testing and cut-over", "4 weeks", "Email, load tests, go-live"],
        ["Hypercare", "3 weeks", "Daily review, tuning, handover"],
    ]
    t = Table(data, colWidths=[70 * mm, 28 * mm, 72 * mm])
    t.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9.5),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF4")),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#999999")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    S.append(t)
    S.append(Spacer(1, 8))
    P("The total elapsed duration is 19 weeks from the start of mobilisation. Dates will be fixed once the "
      "customer confirms the start date and nominates its delivery team.")

    P("6. Commercial Assumptions", h1)
    P("The engagement is priced on a fixed-fee basis for the scope in Section 3, with milestone payments at "
      "the end of each phase in Section 5. Cloud consumption charges are billed directly to the customer's AWS "
      "account and are not included in the fee. Brightpath Cloud Consulting has assumed that the customer "
      "will provide timely access to its AWS account, to its PayFlow sandbox credentials and to subject matter "
      "experts for up to four hours per week. Change requests that alter the scope in Section 3 will be "
      "estimated and agreed in writing before work begins.")
    P("Support after the hypercare period is not included and may be contracted separately under a managed "
      "services agreement. Training for up to eight of the customer's staff is included.")

    P("7. Risks and Dependencies", h1)
    P("The main delivery risk is the availability of the customer's engineering lead for design questions "
      "during the order-processing phase. The main technical dependency is the PayFlow sandbox, which must "
      "support the load tests planned for the cut-over phase. Both items will be reviewed in the weekly "
      "steering meeting and escalated to the sponsor where necessary.")

    doc = SimpleDocTemplate(str(HERE / "b_prose.pdf"), pagesize=A4,
                            leftMargin=20 * mm, rightMargin=20 * mm,
                            topMargin=20 * mm, bottomMargin=20 * mm,
                            title="Solution Design Document - Order Processing Modernisation",
                            author="Brightpath Cloud Consulting (fictional)")
    doc.build(S, onFirstPage=footer, onLaterPages=footer)


# --------------------------------------------------------------------------
# c_tables.docx - architecture mostly in tables
# --------------------------------------------------------------------------
def make_c_tables() -> None:
    from docx import Document
    from docx.shared import Pt

    d = Document()
    style = d.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    d.add_heading("Data Analytics Platform - Solution Architecture", 0)
    d.add_paragraph("Customer: Tailwind Manufacturing (fictional)    Version: 1.0 FINAL    "
                    "Region: eu-central-1")

    d.add_heading("1. Purpose", 1)
    d.add_paragraph(
        "This document records the final architecture of the Tailwind analytics platform. It consolidates "
        "production, quality and sales data from the customer's corporate data center into an AWS data lake "
        "and exposes it to analysts through SQL and dashboards. The components and integrations are "
        "listed in the tables below; the tables are the authoritative description.")

    d.add_heading("2. Component Inventory", 1)
    comps = [
        ("Component", "AWS Service", "Placement", "Purpose"),
        ("Corporate data center", "Not applicable (on-premises)", "Customer premises (on_prem)",
         "Hosts the source systems and the VPN device"),
        ("Source database", "Microsoft SQL Server 2019 (on-premises)", "Corporate data center",
         "System of record for production and sales data"),
        ("VPN device", "Customer gateway (on-premises router)", "Corporate data center",
         "Terminates the IPsec tunnels to AWS"),
        ("Site-to-Site VPN", "AWS Site-to-Site VPN", "Region eu-central-1 / VPC edge",
         "Encrypted connectivity between the data center and the VPC"),
        ("Analytics VPC", "Amazon VPC", "Region eu-central-1",
         "Network boundary for all platform components that run in a VPC"),
        ("Private subnets", "Amazon VPC private subnets (two Availability Zones)", "Analytics VPC",
         "Host the replication instance and interface endpoints; no internet route"),
        ("Replication service", "AWS DMS (AWS Database Migration Service)", "Private subnets",
         "Full load and ongoing change data capture from SQL Server"),
        ("S3 gateway endpoint", "VPC endpoint (gateway) for Amazon S3", "Analytics VPC",
         "Private access from the VPC to Amazon S3"),
        ("Glue interface endpoint", "VPC endpoint (interface) for AWS Glue", "Private subnets",
         "Private access from the VPC to the AWS Glue API"),
        ("Raw zone", "Amazon S3 (raw bucket)", "Region eu-central-1",
         "Landing area for DMS output, Parquet, partitioned by date"),
        ("Curated zone", "Amazon S3 (curated bucket)", "Region eu-central-1",
         "Cleaned and modelled datasets for analytics"),
        ("ETL and catalog", "AWS Glue (ETL jobs, crawlers and Data Catalog)", "Region eu-central-1",
         "Transforms raw to curated and maintains table metadata"),
        ("Data governance", "AWS Lake Formation", "Region eu-central-1",
         "Table and column level permissions on the data lake"),
        ("Ad-hoc query", "Amazon Athena", "Region eu-central-1",
         "Serverless SQL over the curated zone"),
        ("Data warehouse", "Amazon Redshift Serverless", "Region eu-central-1",
         "Governed, high-concurrency reporting workloads"),
        ("Dashboards", "Amazon QuickSight", "Region eu-central-1",
         "Business dashboards for the customer's analysts"),
        ("Encryption", "AWS KMS", "Region eu-central-1",
         "Customer managed keys for S3, Redshift and Glue"),
        ("Monitoring", "Amazon CloudWatch", "Region eu-central-1",
         "Metrics, logs and alarms for DMS and Glue"),
    ]
    t = d.add_table(rows=0, cols=4)
    t.style = "Table Grid"
    for i, row in enumerate(comps):
        cells = t.add_row().cells
        for c, v in zip(cells, row):
            c.text = v
            if i == 0:
                for r in c.paragraphs[0].runs:
                    r.bold = True

    d.add_heading("3. Integrations and Data Flow", 1)
    d.add_paragraph("Each row describes one directed interaction. Steps are in processing order.")
    flows = [
        ("From", "To", "Protocol", "Notes"),
        ("Source database", "VPN device", "TCP 1433 (inside the tunnel)",
         "SQL Server is reachable only on the corporate network"),
        ("Corporate data center", "Site-to-Site VPN", "IPsec",
         "Two redundant tunnels to the virtual private gateway"),
        ("Site-to-Site VPN", "Analytics VPC", "IPsec / private routing",
         "Routes the data center range into the private subnets"),
        ("Source database", "AWS DMS", "TCP 1433 over VPN",
         "DMS pulls full load, then change data capture"),
        ("AWS DMS", "Amazon S3 raw bucket", "HTTPS via S3 gateway endpoint",
         "Writes Parquet files to the raw zone"),
        ("AWS Glue", "Amazon S3 raw bucket", "HTTPS via S3 gateway endpoint",
         "Reads raw files"),
        ("AWS Glue", "Amazon S3 curated bucket", "HTTPS via S3 gateway endpoint",
         "Writes curated datasets, nightly schedule"),
        ("AWS Lake Formation", "Amazon S3 curated bucket", "Internal",
         "Registers the location and enforces permissions"),
        ("Amazon Athena", "AWS Glue Data Catalog", "HTTPS",
         "Resolves table definitions for queries"),
        ("Amazon Athena", "Amazon S3 curated bucket", "HTTPS",
         "Queries curated data in place"),
        ("Amazon Redshift Serverless", "Amazon S3 curated bucket", "HTTPS",
         "Loads and queries curated data"),
        ("Amazon QuickSight", "Amazon Athena", "HTTPS",
         "Dataset for exploratory dashboards"),
        ("Amazon QuickSight", "Amazon Redshift Serverless", "TLS (JDBC)",
         "Dataset for governed reporting dashboards"),
    ]
    t2 = d.add_table(rows=0, cols=4)
    t2.style = "Table Grid"
    for i, row in enumerate(flows):
        cells = t2.add_row().cells
        for c, v in zip(cells, row):
            c.text = v
            if i == 0:
                for r in c.paragraphs[0].runs:
                    r.bold = True

    d.add_heading("4. Notes", 1)
    d.add_paragraph("All data stores are encrypted with customer managed AWS KMS keys. Replication runs "
                    "continuously; the Glue curation jobs run nightly.")
    d.save(str(HERE / "c_tables.docx"))


# --------------------------------------------------------------------------
# e_scanned.pdf - image-only PDF, no text layer
# --------------------------------------------------------------------------
def make_e_scanned() -> None:
    from PIL import Image, ImageDraw, ImageFilter, ImageFont

    lines = [
        ("Claims Intake Platform - Final Architecture", True),
        ("Customer: Harbor Mutual Insurance (fictional)", False),
        ("", False),
        ("Policyholders submit claims through a web portal that is delivered", False),
        ("from Amazon S3. The portal calls Amazon API Gateway, which invokes", False),
        ("AWS Lambda functions. Claim records are stored in Amazon DynamoDB", False),
        ("and supporting photos are stored in Amazon S3. Each new claim is", False),
        ("published to Amazon SNS so that the adjuster notification service", False),
        ("can send an email. Logs go to Amazon CloudWatch.", False),
        ("", False),
        ("Flow: 1. Portal -> API Gateway   2. API Gateway -> Lambda", False),
        ("      3. Lambda -> DynamoDB      4. Lambda -> SNS", False),
    ]

    def load_font(size: int, bold: bool):
        candidates = [
            "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Arial.ttf",
            "/System/Library/Fonts/Helvetica.ttc",
            "/Library/Fonts/Arial.ttf",
        ]
        for c in candidates:
            if os.path.exists(c):
                try:
                    return ImageFont.truetype(c, size)
                except OSError:
                    pass
        return ImageFont.load_default(size)

    W, H = 1654, 2339  # A4 at 200 dpi
    img = Image.new("L", (W, H), 245)
    dr = ImageDraw.Draw(img)
    y = 220
    for text, bold in lines:
        dr.text((140, y), text, fill=25, font=load_font(46 if bold else 38, bold))
        y += 80 if bold else 62

    rng = random.Random(7)  # scanner noise, deterministic
    px = img.load()
    for _ in range(60000):
        x, yy = rng.randrange(W), rng.randrange(H)
        px[x, yy] = max(0, px[x, yy] - rng.randrange(30, 90))
    img = img.rotate(0.6, resample=Image.BICUBIC, fillcolor=245).filter(ImageFilter.GaussianBlur(0.6))
    # Pillow writes an image-only PDF (no text layer)
    img.convert("RGB").save(str(HERE / "e_scanned.pdf"), "PDF", resolution=200.0)


if __name__ == "__main__":
    make_b_prose()
    make_c_tables()
    make_e_scanned()
    print("generated:", ", ".join(sorted(p.name for p in HERE.iterdir() if p.suffix in (".pdf", ".docx"))))
