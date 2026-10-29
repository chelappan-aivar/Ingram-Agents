"""The branded .docx renders from SOW markdown and opens again with python-docx."""

import io

from conftest import PNG_B64, SAMPLE_SOW
from docx import Document

from sow_generator.core.config import CompanyProfile
from sow_generator.schemas import ArchImage
from sow_generator.views.docx import build_sow_docx


def _render(images=None):
    data = build_sow_docx("Acme Corp", SAMPLE_SOW, "Statement of Work", images, company=CompanyProfile())
    assert data[:2] == b"PK"   # a zip container
    return Document(io.BytesIO(data))


def test_docx_has_cover_body_and_acceptance():
    doc = _render()
    texts = [p.text for p in doc.paragraphs]
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert "Statement of Work" in texts and "Acme Corp" in texts               # cover
    assert any("proprietary to Ingram Micro." in t for t in texts)            # important notice
    assert headings[:3] == ["About Ingram Micro", "Statement of Work", "Scope of Work"]
    assert "Acceptance and Authorization" in headings
    assert doc.sections[0].header.paragraphs[0].text == "IngramMicro_Acme_Corp_SOW"
    # Markdown tables in the body render as real docx tables.
    raci = [t for t in doc.tables if t.cell(0, 0).text == "Item"]
    assert raci and raci[0].cell(1, 2).text == "Ingram Micro"
    # The pricing section links to the cost estimate instead of tabulating it.
    assert any("calculator.aws/#/estimate?id=sample123" in t for t in texts)
    assert any(t.cell(0, 1).text == "For Ingram Micro" for t in doc.tables)    # signature table


def test_docx_places_architecture_image():
    without = _render()
    with_image = _render([ArchImage(filename="arch.png", content_b64=f"data:image/png;base64,{PNG_B64}")])
    assert len(with_image.inline_shapes) == len(without.inline_shapes) + 1


# ----------------------------------------------------------------- contents page and formatting
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

FORMAT_SOW = """# Overview

Intro paragraph with **bold** text and a [link](https://example.com/page).

## Pricing

AWS Pricing Calculator estimate: https://calculator.aws/#/estimate?id=5e47657db0735aa49c150cb9646e0dc9683038da

### Details

- top bullet
  - nested bullet
1. first step
2. second step

---

| Item | Owner |
|------|-------|
| A | Ingram |
| B | Customer |
"""


def _format_doc():
    data = build_sow_docx("Acme Corp", FORMAT_SOW, "Statement of Work", None, company=CompanyProfile())
    return Document(io.BytesIO(data))


def test_contents_page_lists_every_heading_before_the_body():
    doc = _format_doc()
    texts = [p.text for p in doc.paragraphs]
    start = texts.index("Contents")
    body_start = next(i for i, p in enumerate(doc.paragraphs) if p.style.name == "Heading 1")
    listed = [t for t in texts[start + 1:body_start] if t.strip()]
    assert listed == ["Overview", "Pricing", "Details", "Acceptance and Authorization"]
    xml = doc.element.xml
    assert 'TOC \\o "1-3"' in xml and xml.count('w:fldCharType="begin"') >= 1
    assert xml.index("TOC \\o") < xml.index('w:fldCharType="end"')
    # Word is asked to fill in the page numbers when the file opens.
    assert doc.settings.element.find(qn("w:updateFields")) is not None


def test_paragraph_with_a_url_is_not_justified():
    doc = _format_doc()
    url = next(p for p in doc.paragraphs if "calculator.aws" in p.text)
    prose = next(p for p in doc.paragraphs if p.text.startswith("Intro paragraph"))
    assert url.alignment != WD_ALIGN_PARAGRAPH.JUSTIFY
    assert prose.alignment != WD_ALIGN_PARAGRAPH.JUSTIFY      # it holds a link too
    plain = build_sow_docx("Acme", "# H\n\nJust ordinary prose that can be justified.\n", company=CompanyProfile())
    para = next(p for p in Document(io.BytesIO(plain)).paragraphs if p.text.startswith("Just ordinary"))
    assert para.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY


def test_urls_and_markdown_links_are_clickable_and_markers_are_gone():
    doc = _format_doc()
    rels = {r.target_ref for r in doc.part.rels.values() if r.is_external}
    assert "https://example.com/page" in rels
    assert "https://calculator.aws/#/estimate?id=5e47657db0735aa49c150cb9646e0dc9683038da" in rels
    texts = " ".join(p.text for p in doc.paragraphs)
    assert "**" not in texts and "](" not in texts and "---" not in texts


def test_lists_keep_their_level_and_numbers():
    doc = _format_doc()
    styles = {p.text: p.style.name for p in doc.paragraphs}
    assert styles["top bullet"] == "List Bullet" and styles["nested bullet"] == "List Bullet 2"
    assert any(p.text.startswith("1.") and "first step" in p.text for p in doc.paragraphs)


def test_table_has_a_repeating_shaded_header():
    doc = _format_doc()
    table = next(t for t in doc.tables if t.cell(0, 0).text == "Item")
    assert table.rows[0]._tr.find(qn("w:trPr")).find(qn("w:tblHeader")) is not None
    fill = table.cell(0, 0)._tc.tcPr.find(qn("w:shd")).get(qn("w:fill"))
    assert fill == "1F3A5F"
    assert table.cell(1, 1).text == "Ingram" or table.cell(1, 0).text == "A"
