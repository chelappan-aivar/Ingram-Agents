"""Shared test fixtures (no AWS). Other work packages may use:

  settings         Settings with a tmp data_dir (no .env)
  store, bus       JobStore / EventBus on that data_dir
  fixtures_dir     Path to tests/fixtures
  fake_llm         FakeLLM (empty queues; set .tool_responses / .text_responses)
  llm_factory      lambda step: fake_llm (same instance for every step)
  stub_stages      factory -> list[(name, fn)] of simple stages, see make_stub_stages()
  make_job         make_job(store, {"name.md": b"..."}) -> Job with input files written
"""
from __future__ import annotations

import io
from pathlib import Path

import pytest

from diagram_generator.config import Settings
from diagram_generator.contracts import HarnessRecord
from diagram_generator.jobs import EventBus, JobStore
from diagram_generator.llm.bedrock import FakeLLM

FIXTURES = Path(__file__).resolve().parent / "fixtures"


def make_settings(tmp_path: Path, **overrides) -> Settings:
    base = dict(DATA_DIR=str(tmp_path / "data"))
    base.update(overrides)
    return Settings(_env_file=None, **base)


# ----------------------------------------------------------------------------- stub stages
def make_stub_stages(fail_at: str | None = None, failure: Exception | None = None, extract: bool = True):
    """Minimal stages that exercise the framework: INGEST reads the files, EXTRACTOR makes one forced
    tool call, HARNESS_1 writes a harness record, PACKAGE writes report.json.
    `fail_at` names the stage that raises `failure` instead of running."""
    from diagram_generator.ingest.extractors import extract_sources
    from diagram_generator.pipeline import StageFailure

    def maybe_fail(ctx, name):
        if fail_at == name:
            raise failure or StageFailure(name, f"stub failure in {name}")

    def ingest(ctx):
        maybe_fail(ctx, "INGEST")
        ctx.emit("INGEST", "started", f"Reading {len(ctx.input_paths)} file(s)")
        sources = extract_sources(ctx.input_paths, ctx.settings.max_text_chars) if extract else []
        ctx.state["sources"] = sources
        ctx.write_json("sources.json", [s.model_dump(exclude={"text"}) for s in sources] or [{"file": "none"}])
        ctx.emit("INGEST", "completed", "Read files", counts={"files": len(sources)})

    def extractor(ctx):
        maybe_fail(ctx, "EXTRACTOR")
        ctx.emit("EXTRACTOR", "started", "Extractor running")
        llm = ctx.llm("extractor")
        out = llm.converse_tool("sys", [{"role": "user", "content": [{"text": "go"}]}],
                                {"toolSpec": {"name": "submit_extraction"}}, ctx.settings.max_tokens_for("extractor"))
        ctx.record_llm_call("extractor", dict(llm.last_call))
        ctx.state["raw"] = out
        ctx.emit("EXTRACTOR", "completed", "Extractor returned", counts={"components": len(out.get("components", []))})

    def harness_1(ctx):
        maybe_fail(ctx, "HARNESS_1")
        ctx.emit("HARNESS_1", "started", "checking")
        n = len(ctx.state["raw"].get("components", []))
        ctx.harness(HarnessRecord(step="harness_1", ok=True, in_counts={"files": len(ctx.input_paths)},
                                  out_counts={"components": n}))

    def package(ctx):
        maybe_fail(ctx, "PACKAGE")
        ctx.write_json("report.json", {"status": "pass"})
        ctx.emit("PACKAGE", "completed", "Report: pass")

    return [("INGEST", ingest), ("EXTRACTOR", extractor), ("HARNESS_1", harness_1), ("PACKAGE", package)]


# ----------------------------------------------------------------------------- fixtures
@pytest.fixture
def settings(tmp_path) -> Settings:
    return make_settings(tmp_path)


@pytest.fixture
def store(settings) -> JobStore:
    return JobStore(settings.jobs_dir)


@pytest.fixture
def bus() -> EventBus:
    return EventBus()


@pytest.fixture
def fixtures_dir() -> Path:
    return FIXTURES


@pytest.fixture
def fake_llm() -> FakeLLM:
    return FakeLLM()


@pytest.fixture
def llm_factory(fake_llm):
    return lambda step: fake_llm


@pytest.fixture
def stub_stages():
    return make_stub_stages


@pytest.fixture
def make_job():
    def _make(store: JobStore, files: dict[str, bytes]):
        job = store.create(list(files))
        for name, data in files.items():
            (store.job_dir(job.job_id) / "input" / name).write_bytes(data)
        return job
    return _make


# ----------------------------------------------------------------------------- document builders
def make_pdf(lines: list[str]) -> bytes:
    """Hand-built single-page text PDF (Helvetica)."""
    content = "BT /F1 11 Tf 40 780 Td 14 TL\n" + "\n".join(
        "(%s) Tj T*" % l.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)") for l in lines) + "\nET"
    objs = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 842] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        "<< /Length %d >>\nstream\n%s\nendstream" % (len(content), content),
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, o in enumerate(objs, 1):
        offsets.append(out.tell())
        out.write(("%d 0 obj\n%s\nendobj\n" % (i, o)).encode("latin-1"))
    xref = out.tell()
    out.write(("xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)).encode())
    for off in offsets:
        out.write(("%010d 00000 n \n" % off).encode())
    out.write(("trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)).encode())
    return out.getvalue()


def make_blank_pdf() -> bytes:
    from pypdf import PdfWriter

    w = PdfWriter()
    w.add_blank_page(width=612, height=792)
    buf = io.BytesIO()
    w.write(buf)
    return buf.getvalue()


def make_docx(path: Path) -> None:
    import docx

    d = docx.Document()
    d.add_heading("Orders architecture", level=1)
    d.add_paragraph("Customers reach an ALB which forwards to ECS.")
    t = d.add_table(rows=2, cols=3)
    for i, h in enumerate(["Component", "Service", "Notes"]):
        t.rows[0].cells[i].text = h
    for i, v in enumerate(["db", "Amazon RDS", "multi-AZ"]):
        t.rows[1].cells[i].text = v
    d.save(str(path))
