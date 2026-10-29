"""Sanity checks for the fixtures.

Run from this folder:
  uv run --python 3.12 --with pypdf --with python-docx python _check.py

Checks: every fixture parses; expectations.json is valid; every must_have item (and both ends of every
must_have connection) appears in the source text; no must_not_have item appears in the source text;
e_scanned.pdf has no text layer; the word "client" never appears in fixture prose.
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
errors: list[str] = []


def norm(s: str) -> str:
    return re.sub(r"\s+", " ", s).lower()


def read_text(path: Path) -> str:
    suf = path.suffix.lower()
    if suf in (".md", ".txt"):
        return path.read_text(encoding="utf-8")
    if suf == ".pdf":
        from pypdf import PdfReader
        r = PdfReader(str(path))
        return "\n".join((p.extract_text() or "") for p in r.pages)
    if suf == ".docx":
        from docx import Document
        d = Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    raise ValueError(path)


def main() -> int:
    exp = json.loads((HERE / "expectations.json").read_text(encoding="utf-8"))
    for fx in exp["fixtures"]:
        f = HERE / fx["file"]
        if not f.exists():
            errors.append(f"{fx['file']}: missing")
            continue
        text = read_text(f)
        nt = norm(text)
        print(f"{fx['file']}: {len(text)} chars extracted")

        if fx["file"] == "e_scanned.pdf":
            from pypdf import PdfReader
            pages = len(PdfReader(str(f)).pages)
            if len(text.strip()) > 20:
                errors.append(f"e_scanned.pdf has a text layer ({len(text.strip())} chars)")
            print(f"  pages={pages}, text chars={len(text.strip())}")
            continue

        if len(nt) < 200:
            errors.append(f"{fx['file']}: too little text")
        for item in fx["must_have_components"]:
            if norm(item) not in nt:
                errors.append(f"{fx['file']}: must_have component not in source: {item!r}")
        for a, b in fx["must_have_connections"]:
            for end in (a, b):
                if norm(end) not in nt:
                    errors.append(f"{fx['file']}: connection end not in source: {end!r}")
        for item in fx["must_not_have_components"]:
            if norm(item) in nt:
                errors.append(f"{fx['file']}: must_not_have item appears in source: {item!r}")
        if re.search(r"\bclients?\b", nt):
            errors.append(f"{fx['file']}: contains the word 'client'")

    if (HERE / "b_prose.pdf").exists():
        from pypdf import PdfReader
        n = len(PdfReader(str(HERE / "b_prose.pdf")).pages)
        print(f"b_prose.pdf pages: {n}")
        if not 3 <= n <= 4:
            errors.append(f"b_prose.pdf has {n} pages (want 3-4)")
    if (HERE / "c_tables.docx").exists():
        from docx import Document
        d = Document(str(HERE / "c_tables.docx"))
        print("c_tables.docx tables:", [(len(t.rows), len(t.columns)) for t in d.tables])
        if len(d.tables) < 2:
            errors.append("c_tables.docx: expected >= 2 tables")

    for name in ("README.md",):
        p = HERE / name
        if p.exists() and re.search(r"\bclients?\b", p.read_text(encoding="utf-8").lower()):
            errors.append(f"{name}: contains the word 'client'")

    if errors:
        print("\nFAILED:")
        for e in errors:
            print(" -", e)
        return 1
    print("\nALL CHECKS PASSED")
    return 0


if __name__ == "__main__":
    sys.exit(main())
