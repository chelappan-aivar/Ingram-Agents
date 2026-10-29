"""Convert the README.md the agent generates into a standard Word document.

Only the Markdown subset the README renderer emits is supported: headings, paragraphs, block quotes,
nested bullets, numbered items, tables, and **bold** / _italic_ / `code` inline text. The Word file therefore
always mirrors the README.
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

_INLINE = re.compile(r"(\*\*.+?\*\*|(?<!\w)_.+?_(?!\w)|`.+?`)")
_BULLET = re.compile(r"^(\s*)[-*] (.*)")
_NUMBER = re.compile(r"^(\s*)\d+\. (.*)")
_CELL_SPLIT = re.compile(r"(?<!\\)\|")
CONTENT_IN = 6.27  # A4 width minus 1 inch margins
COL_IN = {"#": 0.4, "Field": 1.7, "Category": 1.3, "Confidence": 1.0}


def _inline(par, text: str) -> None:
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("_") and part.endswith("_") and len(part) > 2:
            par.add_run(part[1:-1]).italic = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            r = par.add_run(part[1:-1])
            r.font.name = "Consolas"
        else:
            par.add_run(part.replace("\\|", "|"))


def _shade(cell, fill: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def _table(doc, rows: list[list[str]]) -> None:
    header, body = rows[0], rows[1:]
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    fixed = [COL_IN.get(h) for h in header]
    flex = (CONTENT_IN - sum(w for w in fixed if w)) / max(1, fixed.count(None))
    widths = [w or flex for w in fixed]
    for cell, h in zip(t.rows[0].cells, header):
        _shade(cell, "D9D9D9")
        _inline(cell.paragraphs[0], f"**{h}**")
    trpr = t.rows[0]._tr.get_or_add_trPr()
    trpr.append(OxmlElement("w:tblHeader"))
    for row in body:
        cells = t.add_row().cells
        for cell, text in zip(cells, row):
            _inline(cell.paragraphs[0], text)
    for row in t.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = Inches(w)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                for r in p.runs:
                    r.font.size = Pt(9.5)
    doc.add_paragraph()


def _page_number_footer(doc) -> None:
    p = doc.sections[0].footer.paragraphs[0]
    p.add_run("Page ")
    for kind, code in (("begin", None), (None, "PAGE"), ("end", None)):
        r = p.add_run()
        if kind:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            r._r.append(fc)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = code
            r._r.append(it)
    for r in p.runs:
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)


def markdown_to_docx(markdown: str, path: Path) -> None:
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Inches(8.27), Inches(11.69)
    s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Inches(1)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    _page_number_footer(doc)

    lines = markdown.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip() or line.strip() == "---":
            i += 1
            continue
        if line.startswith("|"):  # table: header, separator, rows
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append(lines[i])
                i += 1
            rows = [[c.strip() for c in _CELL_SPLIT.split(b.strip().strip("|"))] for b in block
                    if not re.match(r"^\|\s*:?-{3,}", b)]
            _table(doc, rows)
            continue
        m = re.match(r"^(#{1,3}) (.*)", line)
        if m:
            level = len(m.group(1))
            h = doc.add_heading("", 0 if level == 1 else level - 1)
            _inline(h, m.group(2))
        elif line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            _inline(p, line[2:])
            for r in p.runs:
                r.italic = True
        elif (b := _BULLET.match(line)):
            level = min(len(b.group(1)) // 2, 2)
            p = doc.add_paragraph(style="List Bullet" + ("" if level == 0 else f" {level + 1}"))
            _inline(p, b.group(2))
        elif (n := _NUMBER.match(line)):
            p = doc.add_paragraph(style="List Number")
            _inline(p, n.group(2))
        elif line.startswith("  "):  # continuation line under a list item
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.5)
            _inline(p, line.strip())
        else:
            p = doc.add_paragraph()
            _inline(p, line)
        i += 1

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
