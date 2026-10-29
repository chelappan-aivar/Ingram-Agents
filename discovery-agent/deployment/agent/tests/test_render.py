import copy
import tempfile
import unittest
from pathlib import Path

from docx import Document

from discovery_agent.analysis import pipeline
from discovery_agent.rendering.word import markdown_to_docx
from discovery_agent.rendering.readme import render_markdown, specifics_lines
from tests.helpers import DOCS, GOOD_MERGE, NARRATIVE


def report(**uc_fields):
    r = {"schema_version": "1.0", "meta": pipeline.build_meta(DOCS, "Acme"), **copy.deepcopy(GOOD_MERGE)}
    r["use_cases"][0].update(uc_fields)
    return r


FULL = dict(systems_involved=["Zendesk", "ServiceNow"], current_process=["Form submitted", "Ticket created"],
            data_sources_and_apis=["Wufoo form export"], owners=[{"name": "Justine", "role": "partner support"}],
            dependencies_and_blockers=["No payment APIs"], stated_metrics=["3 days every month"],
            status="planned_by_customer")


class RenderTests(unittest.TestCase):
    def test_specifics_appear_in_fixed_order(self):
        lines = specifics_lines(FULL)
        labels = [l.split(":**")[0].replace("- **", "") for l in lines if l.startswith("- **")]
        self.assertEqual(labels, ["Systems involved", "Current process", "Data sources and APIs", "Owners",
                                  "Dependencies and blockers", "Customer-stated metrics", "Status"])
        self.assertIn("  - Step 2: Ticket created", lines)
        self.assertIn("- **Status:** Planned by the customers", lines)

    def test_empty_and_unknown_fields_are_omitted(self):
        self.assertEqual(specifics_lines({"systems_involved": [], "status": "unknown"}), [])
        self.assertEqual(specifics_lines({}), [])

    def test_readme_has_standard_sections_and_specifics(self):
        md = render_markdown(report(**FULL), NARRATIVE)
        for heading in ("## Executive summary", "## 1. Discovery summary", "## 2. Top hidden opportunities",
                        "## 3. Use cases", "## 4. What we heard", "## 5. Open questions"):
            self.assertIn(heading, md)
        self.assertIn("- **Systems involved:** Zendesk, ServiceNow", md)
        self.assertIn("hypothesis, not confirmed", md)

    def test_resolved_questions_move_to_resolved_section(self):
        r = report()
        r["gaps_and_open_questions"] = [
            {"id": "Q-1", "question": "Who owns approvals?", "reason": "unclear",
             "status": "resolved", "resolved_by": "reviewer", "resolution": "Dana owns them."},
            {"id": "Q-2", "question": "Which figure is right?", "reason": "conflict", "type": "conflict",
             "status": "needs_confirmation", "resolved_by": "agent_assumed", "resolution": "Assumed the later one."},
            {"id": "Q-3", "question": "What is the exception volume?", "reason": "not measured"},  # still open
        ]
        md = render_markdown(r)
        self.assertIn("## 5. Resolved during review", md)
        self.assertIn("## 6. Open questions", md)
        # resolved items appear under Resolved, with their resolution + who resolved
        resolved_block = md.split("## 5. Resolved during review")[1].split("## 6. Open questions")[0]
        self.assertIn("Who owns approvals?", resolved_block)
        self.assertIn("Dana owns them.", resolved_block)
        self.assertIn("answered by reviewer", resolved_block)
        self.assertIn("agent's stand — needs confirmation", resolved_block)
        # the still-open one is NOT in the resolved block, and IS in open questions
        self.assertNotIn("exception volume", resolved_block)
        open_block = md.split("## 6. Open questions")[1]
        self.assertIn("exception volume", open_block)
        self.assertNotIn("Who owns approvals?", open_block)

    def test_evidence_is_limited_to_a_few_quotes(self):
        r = report()
        r["use_cases"][0]["evidence"] = [dict(r["use_cases"][0]["evidence"][0], quote=f"quote number {i}") for i in range(5)]
        md = render_markdown(r)
        self.assertIn("more in the JSON report", md)
        self.assertLessEqual(md.count("quote number"), 2)

    def test_word_document_is_converted_from_the_readme(self):
        md = render_markdown(report(**FULL), NARRATIVE)
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "r.docx"
            markdown_to_docx(md, path)
            doc = Document(str(path))
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("Step 1: Form submitted", text)
        self.assertIn("Systems involved: Zendesk, ServiceNow", text)
        self.assertEqual([p.text for p in doc.paragraphs if p.style.name == "Heading 1"][:2],
                         ["Executive summary", "1. Discovery summary"])
        self.assertGreaterEqual(len(doc.tables), 2)


if __name__ == "__main__":
    unittest.main()
