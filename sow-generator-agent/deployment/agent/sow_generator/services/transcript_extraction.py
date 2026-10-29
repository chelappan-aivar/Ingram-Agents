"""Extract plain text from an uploaded transcript file."""

import base64
import csv
import io
import os

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader


def _strip_data_url(content_b64: str) -> str:
    if content_b64.startswith("data:") and ";base64," in content_b64:
        return content_b64.split(";base64,", 1)[1]
    return content_b64


def _ext(name: str) -> str:
    return os.path.splitext(name)[1].lower().lstrip(".")


def _extract_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    pages: list[str] = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            pages.append("")
    return "\n\n".join(p.strip() for p in pages if p.strip())


def _extract_docx(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts: list[str] = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_text(data: bytes) -> str:
    for enc in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _extract_csv(data: bytes) -> str:
    text = _extract_text(data)
    reader = csv.reader(io.StringIO(text))
    rows = [
        " | ".join(cell.strip() for cell in row) for row in reader if any(c.strip() for c in row)
    ]
    return "\n".join(rows)


def _extract_xlsx(data: bytes) -> str:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    for sheet in wb.worksheets:
        parts.append(f"## Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_transcript_from_file(name: str, content_b64: str) -> str:
    """Decode content_b64 and extract readable text based on file extension."""
    if not content_b64:
        raise ValueError(f"Empty file content for {name!r}")

    raw = base64.b64decode(_strip_data_url(content_b64), validate=False)
    ext = _ext(name)

    if ext == "pdf":
        text = _extract_pdf(raw)
    elif ext == "docx":
        text = _extract_docx(raw)
    elif ext in ("txt", "md", "markdown", "text"):
        text = _extract_text(raw)
    elif ext == "csv":
        text = _extract_csv(raw)
    elif ext == "xlsx":
        text = _extract_xlsx(raw)
    else:
        raise ValueError(
            f"Unsupported file type {ext!r} for {name!r}. Supported: pdf, docx, txt, md, csv, xlsx."
        )

    text = (text or "").strip()
    if not text:
        raise ValueError(f"No readable text extracted from {name!r}")

    header = f"# Attached file: {name}\n\n"
    return header + text
