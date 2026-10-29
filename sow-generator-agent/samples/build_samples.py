#!/usr/bin/env python3
"""Build the SOW generator sample inputs (fictional customers) from the text in this file.

    python3 build_samples.py

Writes, per scenario folder: project-details.txt, upstream.json, transcripts / notes (.docx, .pdf, .txt, .xlsx)
and architecture-diagram.png. Needs python-docx, openpyxl, reportlab, Pillow (and jsonschema to validate
upstream.json against ../resources/contracts/upstream_inputs.v1.schema.json).
All names, companies and e-mail addresses are fictional (example.com).
"""
import json
from pathlib import Path

import docx
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

HERE = Path(__file__).resolve().parent
SCHEMA = HERE.parent / "resources" / "contracts" / "upstream_inputs.v1.schema.json"


# ====================================================================== writers
def write_text(path: Path, text: str) -> None:
    path.write_text(text.strip() + "\n", encoding="utf-8")


def write_docx(path: Path, title: str, meta: str, blocks: list[tuple[str, str]]) -> None:
    """blocks: (kind, text) with kind in h (heading), p (paragraph), b (bullet), s (speaker line 'Name: text')."""
    d = docx.Document()
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(11)
    d.add_heading(title, level=1)
    d.add_paragraph(meta).runs[0].italic = True
    for kind, text in blocks:
        if kind == "h":
            d.add_heading(text, level=2)
        elif kind == "b":
            d.add_paragraph(text, style="List Bullet")
        elif kind == "s":
            name, _, said = text.partition(":")
            para = d.add_paragraph()
            para.add_run(f"{name}:").bold = True
            para.add_run(said)
        else:
            d.add_paragraph(text)
    d.core_properties.author = "Ingram Micro (sample data)"
    d.save(path)


def write_pdf(path: Path, title: str, meta: str, blocks: list[tuple[str, str]]) -> None:
    styles = getSampleStyleSheet()
    story = [Paragraph(title, styles["Title"]), Paragraph(f"<i>{meta}</i>", styles["Normal"]), Spacer(1, 10)]
    for kind, text in blocks:
        if kind == "h":
            story += [Spacer(1, 6), Paragraph(text, styles["Heading2"])]
        elif kind == "b":
            story.append(Paragraph(text, styles["Normal"], bulletText="\u2022"))
        else:
            story.append(Paragraph(text, styles["Normal"]))
        story.append(Spacer(1, 4))
    SimpleDocTemplate(str(path), pagesize=A4, title=title, author="Ingram Micro (sample data)").build(story)


def write_xlsx(path: Path, sheets: dict[str, list[list]]) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    head_fill = PatternFill("solid", fgColor="1F3864")
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = head_fill
        for col in ws.columns:
            width = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(60, max(10, width + 2))
    wb.save(path)


def _font(size: int, bold: bool = False):
    for name in (["Arial Bold.ttf", "Helvetica-Bold.ttf", "DejaVuSans-Bold.ttf"] if bold
                 else ["Arial.ttf", "Helvetica.ttf", "DejaVuSans.ttf"]):
        for root in ("/System/Library/Fonts/Supplemental/", "/Library/Fonts/", "/usr/share/fonts/truetype/dejavu/", ""):
            try:
                return ImageFont.truetype(root + name, size)
            except OSError:
                continue
    return ImageFont.load_default()


def write_diagram(path: Path, title: str, groups: list[dict], arrows: list[tuple[str, str, str]]) -> None:
    """groups: {label, x, y, w, h, color, boxes: [(id, title, subtitle, x, y)]}; arrows: (from_id, to_id, label)."""
    W, H = 1800, 1000
    img = Image.new("RGB", (W, H), "white")
    g = ImageDraw.Draw(img)
    f_title, f_group, f_box, f_sub, f_arrow = _font(34, True), _font(22, True), _font(20, True), _font(16), _font(15)
    g.text((40, 26), title, font=f_title, fill="#1F3864")
    centres = {}
    BW, BH = 230, 86
    for grp in groups:
        x, y, w, h = grp["x"], grp["y"], grp["w"], grp["h"]
        g.rounded_rectangle((x, y, x + w, y + h), radius=16, outline=grp["color"], width=3, fill="#FAFBFD")
        g.text((x + 16, y + 10), grp["label"], font=f_group, fill=grp["color"])
        for bid, btitle, bsub, bx, by in grp["boxes"]:
            g.rounded_rectangle((bx, by, bx + BW, by + BH), radius=10, outline="#5B6B82", width=2, fill="white")
            g.text((bx + BW / 2, by + 30), btitle, font=f_box, fill="#111", anchor="mm")
            g.text((bx + BW / 2, by + 60), bsub, font=f_sub, fill="#555", anchor="mm")
            centres[bid] = (bx, by, bx + BW, by + BH)
    for a, b, label in arrows:
        ax0, ay0, ax1, ay1 = centres[a]
        bx0, by0, bx1, by1 = centres[b]
        if bx0 >= ax1:            # left -> right
            p, q = (ax1, (ay0 + ay1) / 2), (bx0, (by0 + by1) / 2)
        elif bx1 <= ax0:          # right -> left
            p, q = (ax0, (ay0 + ay1) / 2), (bx1, (by0 + by1) / 2)
        elif by0 >= ay1:          # top -> bottom
            p, q = ((ax0 + ax1) / 2, ay1), ((bx0 + bx1) / 2, by0)
        else:                     # bottom -> top
            p, q = ((ax0 + ax1) / 2, ay0), ((bx0 + bx1) / 2, by1)
        g.line((p, q), fill="#E06C00", width=3)
        vx, vy = q[0] - p[0], q[1] - p[1]
        n = max(1.0, (vx * vx + vy * vy) ** 0.5)
        ux, uy = vx / n, vy / n
        g.polygon([q, (q[0] - 14 * ux + 7 * uy, q[1] - 14 * uy - 7 * ux), (q[0] - 14 * ux - 7 * uy, q[1] - 14 * uy + 7 * ux)],
                  fill="#E06C00")
        if label:
            g.text(((p[0] + q[0]) / 2, (p[1] + q[1]) / 2 - 12), label, font=f_arrow, fill="#7A3B00", anchor="mm")
    g.text((W - 40, H - 24), "Sample data - fictional customer", font=f_sub, fill="#999", anchor="rs")
    img.save(path, optimize=True)


def validate(upstream: dict, where: Path) -> None:
    try:
        import jsonschema
    except ImportError:
        print(f"  (jsonschema not installed; {where.name} not validated)")
        return
    jsonschema.validate(upstream, json.loads(SCHEMA.read_text(encoding="utf-8")))
    cost = upstream.get("cost") or {}
    items = cost.get("line_items") or []
    if items:
        monthly = round(sum(i.get("monthly") or 0 for i in items), 2)
        assert monthly == cost["monthly_total"], f"{where}: line items add up to {monthly}, not {cost['monthly_total']}"
        assert round(monthly * 12, 2) == cost["annual_total"], f"{where}: annual_total must be 12 x monthly_total"


# ====================================================================== 01: AI/ML - Fabrikam Insurance
def scenario_aiml(out: Path) -> None:
    write_text(out / "project-details.txt", """
Fabrikam Insurance - Claims FNOL Assistant (generative AI)

Customer: Fabrikam Insurance, a US property and casualty insurer (home and auto), about 2,100 employees, HQ in Columbus, Ohio.
Opportunity: OPP-2026-0517. Ingram Micro seller: Priya Nair. Template: AI/ML.

Problem
Fabrikam receives about 18,000 first notice of loss (FNOL) reports per month by phone, e-mail and a web form. Claims
intake agents re-key each report into Guidewire ClaimCenter, read attached police reports, repair estimates and photos,
and check the policy for coverage. Average handling time is 22 minutes per FNOL and 9% of new claims need rework
because of missing or wrong data. Peak volumes after storms reach 3x normal and create a 2-day backlog.

What they want
- A generative AI assistant that reads FNOL e-mails, web-form submissions and attachments (PDF, images, Word).
- Extract claimant, policy number, loss date, location, cause of loss, vehicles / property, injuries and third parties.
- Check the policy and coverage using the policy wording documents (about 1,400 PDFs) with retrieval (RAG).
- Pre-fill a draft claim in Guidewire ClaimCenter through its REST API; an intake agent reviews and approves it.
- Flag missing information and draft a follow-up e-mail to the claimant (English and Spanish).
- Detect and redact PII in logs; nothing leaves the US (us-east-2, Ohio).
- Dashboard for intake supervisors: volume, handling time, extraction accuracy, approval rate.

Targets
- Average handling time from 22 minutes to under 8 minutes.
- Rework rate from 9% to under 3%.
- Field extraction accuracy of at least 95% on the agreed test set of 500 historical FNOLs.

Constraints
- Human in the loop: no claim is created without intake-agent approval.
- Data residency: US only (us-east-2). SOC 2 Type II controls; no customer data used to train models.
- Sign-in with Fabrikam's Microsoft Entra ID (SSO) for intake agents and supervisors.

Timeline and budget
- 10 weeks, starting 11 January 2027; pilot with the auto claims team (40 intake agents), then home claims.
- One-time delivery budget around USD 190,000. Expected AWS run-rate about USD 8,500 per month.
""")

    upstream = {
        "meta": {
            "schema_version": "1.0",
            "opportunity_id": "OPP-2026-0517",
            "template_id": "aiml/ai-ml",
            "source_agents": ["intake", "discovery", "architecture", "cost-estimate"],
            "created_at": "2026-10-22T10:30:00Z",
            "generated_by": "sample-data",
        },
        "intake": {
            "customer_name": "Fabrikam Insurance",
            "industry": "Insurance (property and casualty)",
            "contacts": [
                {"name": "Laura Chen", "role": "VP Claims Operations (executive sponsor)", "email": "laura.chen@fabrikam-insurance.example.com"},
                {"name": "Marcus Reid", "role": "Director, Enterprise Architecture", "email": "marcus.reid@fabrikam-insurance.example.com"},
                {"name": "Elena Torres", "role": "Claims Intake Manager", "email": "elena.torres@fabrikam-insurance.example.com"},
            ],
            "engagement_type": "Generative AI solution build (pilot to production)",
            "template_hint": "aiml/ai-ml",
            "problem_statement": (
                "Claims intake agents re-key about 18,000 first notice of loss (FNOL) reports per month from phone, e-mail "
                "and web forms into Guidewire ClaimCenter. Average handling time is 22 minutes per FNOL, 9% of new claims "
                "need rework, and storm peaks (3x volume) create a 2-day backlog."
            ),
            "goals": [
                "Reduce average FNOL handling time from 22 minutes to under 8 minutes",
                "Reduce claim rework from 9% to under 3%",
                "Reach at least 95% field extraction accuracy on 500 historical FNOLs",
                "Absorb 3x storm-peak volumes without a backlog",
            ],
            "timeline": "10 weeks",
            "budget_range": "USD 180,000 - 200,000 one-time",
            "constraints": [
                "Human approval before any claim is created in ClaimCenter",
                "All data stays in the US (us-east-2)",
                "SOC 2 Type II controls; no customer data used for model training",
                "Single sign-on with Microsoft Entra ID",
            ],
        },
        "discovery": {
            "source": "discovery",
            "findings": [
                "FNOLs arrive 55% by phone (transcribed by the contact centre), 30% by e-mail and 15% by web form",
                "Attachments: police reports (PDF), repair estimates (PDF/Word), up to 10 photos per claim",
                "Guidewire ClaimCenter 10 exposes REST APIs for draft claim creation (Cloud API)",
                "Policy wording: about 1,400 PDF documents across 6 product lines, updated quarterly",
                "Spanish is used in about 18% of FNOLs",
            ],
            "requirements": [
                "Extract 32 FNOL fields with a confidence score per field",
                "Coverage check with citations to the policy wording",
                "Draft follow-up e-mail for missing information (English and Spanish)",
                "Supervisor dashboard: volume, handling time, extraction accuracy, approval rate",
                "PII redaction in application logs",
            ],
            "stakeholders": [
                {"name": "Laura Chen", "role": "VP Claims Operations"},
                {"name": "Marcus Reid", "role": "Director, Enterprise Architecture"},
                {"name": "Elena Torres", "role": "Claims Intake Manager"},
                {"name": "Sam Patel", "role": "Information Security Officer"},
            ],
            "transcript_refs": [],
            "notes": "Pilot with the auto claims team (40 intake agents) before home claims.",
        },
        "architecture": {
            "diagram_images": [],
            "services": [
                {"name": "Amazon Bedrock", "purpose": "Claude models for extraction, coverage reasoning and e-mail drafts"},
                {"name": "Amazon Bedrock AgentCore", "purpose": "Runtime for the FNOL agent"},
                {"name": "Amazon Bedrock Knowledge Bases", "purpose": "RAG over the policy wording documents"},
                {"name": "Amazon OpenSearch Serverless", "purpose": "Vector store for the knowledge base"},
                {"name": "Amazon Textract", "purpose": "Text and form extraction from PDF and image attachments"},
                {"name": "Amazon Comprehend", "purpose": "PII detection and redaction in logs"},
                {"name": "AWS Lambda", "purpose": "Intake, orchestration and ClaimCenter integration functions"},
                {"name": "Amazon API Gateway", "purpose": "API for the review UI and the web form"},
                {"name": "Amazon S3", "purpose": "Attachments and policy document storage"},
                {"name": "Amazon DynamoDB", "purpose": "FNOL state, review decisions and audit trail"},
                {"name": "Amazon Cognito", "purpose": "Federated sign-in with Microsoft Entra ID"},
                {"name": "Amazon CloudWatch", "purpose": "Logs, metrics, alarms and the supervisor dashboard"},
            ],
            "assumptions": [
                "Fabrikam provides a ClaimCenter sandbox with Cloud API access by week 2",
                "The contact centre already produces text transcripts of phone FNOLs",
                "500 labelled historical FNOLs are available for the test set",
            ],
            "nfrs": [
                "Draft claim ready in under 2 minutes for 95% of FNOLs",
                "Handle 1,800 FNOLs per day at storm peak",
                "99.9% availability during business hours",
                "Encryption at rest (KMS) and in transit (TLS 1.2+)",
            ],
        },
        "cost": {
            "currency": "USD",
            "monthly_total": 8500,
            "annual_total": 102000,
            "region": "us-east-2",
            "line_items": [
                {"service": "Amazon Bedrock (Claude)", "monthly": 4200, "annual": 50400, "notes": "~18,000 FNOLs, ~25k tokens each"},
                {"service": "Amazon Bedrock AgentCore Runtime", "monthly": 650, "annual": 7800, "notes": "Agent sessions"},
                {"service": "Amazon OpenSearch Serverless", "monthly": 1400, "annual": 16800, "notes": "Vector store, 2 OCU minimum"},
                {"service": "Amazon Textract", "monthly": 1100, "annual": 13200, "notes": "~70,000 pages per month"},
                {"service": "Amazon Comprehend", "monthly": 300, "annual": 3600, "notes": "PII detection"},
                {"service": "AWS Lambda and Amazon API Gateway", "monthly": 250, "annual": 3000, "notes": ""},
                {"service": "Amazon S3", "monthly": 120, "annual": 1440, "notes": "Attachments, 1-year retention"},
                {"service": "Amazon DynamoDB", "monthly": 180, "annual": 2160, "notes": "On-demand"},
                {"service": "Amazon CloudWatch", "monthly": 200, "annual": 2400, "notes": "Logs, metrics, dashboards"},
                {"service": "Amazon Cognito", "monthly": 100, "annual": 1200, "notes": "Federated users"},
            ],
            "assumptions": [
                "On-demand pricing in us-east-2, October 2026",
                "18,000 FNOLs per month on average, 3x for storm peaks a few days per month",
            ],
            # calculator_url from the cost estimate agent. The SOW refers to this link
            # instead of restating the run-rate (see pricing-refers-to-cost-estimate).
            "source_link": "https://calculator.aws/#/estimate?id=sample-fabrikam-fnol",
            "source_file": None,
        },
    }
    validate(upstream, out / "upstream.json")
    (out / "upstream.json").write_text(json.dumps(upstream, indent=2) + "\n", encoding="utf-8")

    write_docx(out / "discovery-call-transcript.docx", "Discovery call transcript: Fabrikam Insurance FNOL assistant",
               "21 October 2026, 60 minutes, Microsoft Teams. Ingram Micro: Priya Nair (seller), Arjun Rao (AI/ML architect). "
               "Fabrikam: Laura Chen, Marcus Reid, Elena Torres, Sam Patel.", [
        ("s", "Priya Nair: Thanks for joining. Laura, could you start with what is not working in claims intake today?"),
        ("s", "Laura Chen: We get about eighteen thousand first notices of loss a month. Every one of them is typed into ClaimCenter by hand. It takes twenty-two minutes on average, and after a big storm we are two days behind."),
        ("s", "Elena Torres: And about nine percent come back for rework. Wrong policy number, missing loss location, the third party's insurer is not captured. That is where most of the time goes."),
        ("s", "Arjun Rao: Which channels do the reports come from?"),
        ("s", "Elena Torres: Fifty-five percent phone, but the contact centre already transcribes those calls. Thirty percent e-mail, fifteen percent our web form. E-mails usually have attachments: police reports, repair estimates, photos. Sometimes ten photos."),
        ("s", "Arjun Rao: So the assistant reads the transcript or e-mail and the attachments, extracts the fields and pre-fills a draft claim. Does ClaimCenter give us an API for that?"),
        ("s", "Marcus Reid: Yes, we are on ClaimCenter 10 with the Cloud API. We can give you a sandbox in the second week of the project."),
        ("s", "Arjun Rao: How many fields do you need on the draft?"),
        ("s", "Elena Torres: Thirty-two. I will send the list. The important ones are policy number, loss date and time, location, cause of loss, vehicles or property, injuries and third parties."),
        ("s", "Laura Chen: We also want it to check coverage. Is this loss covered under the policy, and why. Our agents look that up in the policy wording today."),
        ("s", "Marcus Reid: That is about fourteen hundred PDFs across six product lines. They change every quarter."),
        ("s", "Arjun Rao: We would index those in a Bedrock knowledge base and have the agent cite the clause it relied on. The intake agent sees the citation and decides."),
        ("s", "Sam Patel: From a security point of view: nothing leaves the US, we use Ohio today. No training on our data. And logs must not contain claimant PII."),
        ("s", "Arjun Rao: Understood. us-east-2, Bedrock does not train on customer prompts, and we would redact PII in the logs with Comprehend."),
        ("s", "Laura Chen: One more thing. Nobody gets a claim created by a robot. A person approves every draft."),
        ("s", "Priya Nair: That is the design: the assistant proposes, the intake agent approves. What about language?"),
        ("s", "Elena Torres: About eighteen percent are in Spanish. The follow-up e-mails for missing information should go out in the claimant's language."),
        ("s", "Priya Nair: How would you measure success?"),
        ("s", "Laura Chen: Handling time under eight minutes, rework under three percent. And it has to keep up with storm peaks, three times normal volume."),
        ("s", "Arjun Rao: For accuracy we suggest a test set of five hundred historical FNOLs with the correct fields, and a target of ninety-five percent field accuracy."),
        ("s", "Elena Torres: We can label those. Start with auto claims, my forty agents there would be the pilot."),
        ("s", "Priya Nair: On timing and budget?"),
        ("s", "Laura Chen: We would like to start on the eleventh of January and be live with the auto pilot in about ten weeks. We have roughly one hundred and ninety thousand dollars approved for delivery."),
        ("s", "Marcus Reid: Sign-in has to be with our Entra ID. And the supervisors want a dashboard: volume, handling time, accuracy, approval rate."),
        ("s", "Priya Nair: Great. We will send the statement of work next week."),
    ])

    write_pdf(out / "claims-intake-process-notes.pdf", "Claims intake: current process and FNOL field list",
              "Prepared by Elena Torres (Claims Intake Manager), Fabrikam Insurance, 22 October 2026. Sample data.", [
        ("h", "Current process (as-is)"),
        ("b", "1. FNOL received by phone (transcribed), e-mail or web form; e-mails are triaged from a shared mailbox."),
        ("b", "2. Intake agent searches the policy in the policy admin system and opens the policy wording PDF."),
        ("b", "3. Agent reads attachments and types 32 fields into ClaimCenter (about 14 of the 22 minutes)."),
        ("b", "4. Coverage note written by hand; complex cases referred to a senior adjuster."),
        ("b", "5. Missing information requested by e-mail from a template (English only today)."),
        ("h", "Pain points"),
        ("b", "9% of claims need rework; the top causes are wrong policy number (31%), missing loss location (24%) and missing third-party details (19%)."),
        ("b", "Storm peaks: up to 1,800 FNOLs per day, a 2-day backlog and overtime for 3 weeks."),
        ("b", "Spanish-speaking claimants (18%) wait longer for follow-ups."),
        ("h", "Key FNOL fields (12 of 32)"),
        ("b", "Policy number, product line, insured name, claimant name and contact, preferred language"),
        ("b", "Loss date and time, loss location (address and geo-code), cause of loss"),
        ("b", "Vehicles (VIN, make, model, damage) or property (address, damaged areas)"),
        ("b", "Injuries (yes / no, description), police report number, third parties and their insurer"),
        ("h", "Pilot scope"),
        ("p", "Auto claims, 40 intake agents, e-mail and web-form FNOLs plus phone transcripts. Home claims follow after "
              "the pilot is accepted. Out of scope for the pilot: commercial lines, fraud scoring, payments."),
    ])

    write_diagram(out / "architecture-diagram.png", "Fabrikam Insurance - FNOL Assistant on AWS (us-east-2)", [
        {"label": "Channels", "x": 30, "y": 100, "w": 290, "h": 520, "color": "#2E75B6", "boxes": [
            ("mail", "FNOL e-mail", "shared mailbox", 60, 160), ("web", "Web form", "fabrikam.example.com", 60, 300),
            ("ccc", "Phone transcripts", "contact centre", 60, 440)]},
        {"label": "AWS - FNOL assistant", "x": 370, "y": 100, "w": 1050, "h": 700, "color": "#E06C00", "boxes": [
            ("cognito", "Cognito", "Entra ID SSO", 395, 160), ("s3", "Amazon S3", "attachments, policies", 655, 160),
            ("textract", "Amazon Textract", "attachments", 915, 160), ("bedrock", "Amazon Bedrock", "Claude", 1175, 160),
            ("apigw", "API Gateway", "intake API", 395, 330), ("lambda", "AWS Lambda", "intake + orchestration", 655, 330),
            ("agent", "AgentCore Runtime", "FNOL agent", 915, 330),
            ("cw", "CloudWatch", "logs + dashboard", 395, 520), ("ddb", "DynamoDB", "state + audit", 655, 520),
            ("comprehend", "Comprehend", "PII redaction", 915, 520), ("kb", "Knowledge Bases", "OpenSearch Serverless", 1175, 520)]},
        {"label": "Fabrikam", "x": 1470, "y": 100, "w": 300, "h": 700, "color": "#548235", "boxes": [
            ("ui", "Review UI", "40 intake agents", 1505, 330), ("gw", "ClaimCenter", "Cloud API (draft)", 1505, 540)]},
    ], [("mail", "apigw", ""), ("web", "apigw", ""), ("ccc", "apigw", ""), ("cognito", "apigw", ""),
        ("apigw", "lambda", ""), ("lambda", "agent", "invoke"), ("lambda", "s3", ""), ("lambda", "ddb", ""),
        ("lambda", "cw", ""), ("agent", "textract", ""), ("agent", "bedrock", ""), ("agent", "comprehend", ""),
        ("agent", "kb", "RAG"), ("agent", "ui", "draft claim"), ("ui", "gw", "approve")])


# ====================================================================== 02: DevOps - Tailspin Toys
def scenario_devops(out: Path) -> None:
    write_text(out / "project-details.txt", """
Tailspin Toys - DevOps platform for the e-commerce site (CI/CD, containers, observability)

Customer: Tailspin Toys, a toy manufacturer and online retailer in Frankfurt, Germany, about 1,200 employees.
Opportunity: OPP-2026-0498. Ingram Micro seller: Priya Nair. Template: DevOps.
Context: Tailspin is moving its ERP, e-commerce platform and file servers from a Frankfurt colocation facility to AWS
eu-central-1 (separate migration project). This engagement modernises how the e-commerce platform is built and run.

Today
- E-commerce: 9 .NET services (6 on .NET 8, 3 on .NET Framework 4.8), a React storefront and SQL Server.
- Code in Azure DevOps Repos; builds on one Jenkins server; deployments are manual scripts run by 2 people.
- Releases once a month in a 6-hour Saturday maintenance window; 2 to 3 failed releases per quarter.
- No central logging; incidents found by customers. MTTR about 5 hours.
- Christmas season traffic peaks at about 5x normal (mid-November to 24 December).

What they want
- Infrastructure as code with Terraform for all environments (dev, test, prod) in eu-central-1.
- Move source control and CI/CD to GitHub Enterprise Cloud with GitHub Actions; reusable pipeline templates.
- Containerise the 6 .NET 8 services and the storefront; run them on Amazon EKS with Karpenter autoscaling.
- Blue/green or canary deployments with automatic rollback; no maintenance window.
- Observability: CloudWatch Container Insights, Amazon Managed Service for Prometheus and Amazon Managed Grafana,
  alerts to Microsoft Teams and PagerDuty.
- Security in the pipeline: image scanning (Amazon Inspector / ECR), SAST and dependency scanning, AWS Security Hub,
  secrets in AWS Secrets Manager.
- Load test proving 5x peak before the Christmas code freeze.

Targets (DORA)
- Deployment frequency: from monthly to at least weekly per service (daily possible).
- Lead time for changes: under 1 day. Change failure rate: under 10%. MTTR: under 1 hour.

Constraints
- GDPR: all data and logs stay in the EU (eu-central-1).
- Code freeze from 1 November 2027; the platform must be live and load-tested before that.
- The 3 .NET Framework 4.8 services stay on Windows EC2 (migration project); only their pipelines move.
- Tailspin's platform team (4 engineers) must be able to run the platform after handover.

Timeline and budget
- 12 weeks, starting 1 March 2027. Delivery budget about EUR 150,000 (shown in USD in the SOW: about USD 165,000).
- Expected AWS run-rate for the platform about USD 5,800 per month.
""")

    upstream = {
        "meta": {
            "schema_version": "1.0",
            "opportunity_id": "OPP-2026-0498",
            "template_id": "devops/devops-standard",
            "source_agents": ["intake", "discovery", "architecture", "cost-estimate"],
            "created_at": "2026-10-23T09:00:00Z",
            "generated_by": "sample-data",
        },
        "intake": {
            "customer_name": "Tailspin Toys",
            "industry": "Retail and manufacturing (toys, e-commerce)",
            "contacts": [
                {"name": "Daniel Weber", "role": "CTO (executive sponsor)", "email": "daniel.weber@tailspin-toys.example.com"},
                {"name": "Sophie Klein", "role": "IT Operations Lead", "email": "sophie.klein@tailspin-toys.example.com"},
                {"name": "Jonas Becker", "role": "Lead Developer, E-commerce", "email": "jonas.becker@tailspin-toys.example.com"},
            ],
            "engagement_type": "DevOps platform implementation",
            "template_hint": "devops/devops-standard",
            "problem_statement": (
                "The e-commerce platform is released once a month in a 6-hour maintenance window with manual scripts, "
                "2 to 3 releases per quarter fail, there is no central logging and MTTR is about 5 hours. Traffic peaks at "
                "5x during the Christmas season."
            ),
            "goals": [
                "Deploy each service at least weekly without a maintenance window",
                "Lead time for changes under 1 day",
                "Change failure rate under 10% and MTTR under 1 hour",
                "Prove 5x Christmas peak capacity with a load test before the code freeze",
            ],
            "timeline": "12 weeks",
            "budget_range": "EUR 150,000 (about USD 165,000) one-time",
            "constraints": [
                "GDPR: all data and logs stay in eu-central-1",
                "Live and load-tested before the 1 November 2027 code freeze",
                "3 .NET Framework 4.8 services stay on Windows EC2; only their pipelines move",
                "Tailspin's 4-person platform team must run the platform after handover",
            ],
        },
        "discovery": {
            "source": "discovery",
            "findings": [
                "9 .NET services (6 on .NET 8, 3 on .NET Framework 4.8), React storefront, SQL Server",
                "Azure DevOps Repos and one Jenkins server; deployments by manual PowerShell scripts",
                "No infrastructure as code; environments differ (config drift)",
                "Licences for GitHub Enterprise Cloud already purchased for 2027",
                "PagerDuty and Microsoft Teams are used for on-call and alerts",
            ],
            "requirements": [
                "Terraform modules for networking, EKS, ECR, IAM, observability and security",
                "Reusable GitHub Actions workflows with build, test, scan, deploy stages",
                "Canary deployments with automatic rollback on SLO breach",
                "Dashboards and alerts per service; runbooks for the top 10 incidents",
                "Load test at 5x peak (about 2,500 requests per second)",
            ],
            "stakeholders": [
                {"name": "Daniel Weber", "role": "CTO"},
                {"name": "Sophie Klein", "role": "IT Operations Lead"},
                {"name": "Jonas Becker", "role": "Lead Developer, E-commerce"},
            ],
            "transcript_refs": [],
            "notes": "The migration project moves SQL Server to Amazon RDS for SQL Server; this engagement does not.",
        },
        "architecture": {
            "diagram_images": [],
            "services": [
                {"name": "Amazon EKS", "purpose": "Kubernetes clusters for the e-commerce services (non-prod and prod)"},
                {"name": "Karpenter", "purpose": "Node autoscaling for the 5x seasonal peak"},
                {"name": "Amazon ECR", "purpose": "Container image registry with scanning"},
                {"name": "GitHub Actions", "purpose": "CI/CD pipelines (GitHub Enterprise Cloud)"},
                {"name": "Terraform", "purpose": "Infrastructure as code for all environments"},
                {"name": "Argo Rollouts", "purpose": "Canary and blue/green deployments with automatic rollback"},
                {"name": "Amazon Managed Service for Prometheus", "purpose": "Metrics"},
                {"name": "Amazon Managed Grafana", "purpose": "Dashboards"},
                {"name": "Amazon CloudWatch", "purpose": "Container Insights, logs and alarms"},
                {"name": "Amazon Inspector", "purpose": "Container image and EC2 vulnerability scanning"},
                {"name": "AWS Security Hub", "purpose": "Security findings and standards"},
                {"name": "AWS Secrets Manager", "purpose": "Application secrets"},
                {"name": "Application Load Balancer", "purpose": "Ingress for the storefront and APIs"},
            ],
            "assumptions": [
                "The landing zone and VPCs in eu-central-1 are delivered by the migration project by week 2",
                "GitHub Enterprise Cloud organisation is available with SSO by week 1",
                "Tailspin developers are available to fix application issues found during containerisation",
            ],
            "nfrs": [
                "Handle 5x normal traffic (about 2,500 requests per second) with p95 latency under 400 ms",
                "Zero-downtime deployments",
                "99.9% availability for the storefront",
                "All data and logs in eu-central-1 (GDPR)",
            ],
        },
        "cost": {
            "currency": "USD",
            "monthly_total": 5800,
            "annual_total": 69600,
            "region": "eu-central-1",
            "line_items": [
                {"service": "Amazon EKS control plane (2 clusters)", "monthly": 146, "annual": 1752, "notes": "non-prod and prod"},
                {"service": "Amazon EC2 worker nodes (Karpenter)", "monthly": 3904, "annual": 46848, "notes": "Average incl. seasonal peak, Savings Plan"},
                {"service": "Amazon ECR", "monthly": 60, "annual": 720, "notes": "Images and scanning"},
                {"service": "Amazon Managed Service for Prometheus", "monthly": 350, "annual": 4200, "notes": ""},
                {"service": "Amazon Managed Grafana", "monthly": 90, "annual": 1080, "notes": "10 editor/viewer users"},
                {"service": "Amazon CloudWatch", "monthly": 420, "annual": 5040, "notes": "Container Insights, logs, alarms"},
                {"service": "Amazon Inspector", "monthly": 180, "annual": 2160, "notes": ""},
                {"service": "AWS Security Hub", "monthly": 120, "annual": 1440, "notes": ""},
                {"service": "Application Load Balancer", "monthly": 260, "annual": 3120, "notes": "2 ALBs"},
                {"service": "NAT gateways", "monthly": 270, "annual": 3240, "notes": "3 AZs, prod"},
            ],
            "assumptions": [
                "On-demand pricing in eu-central-1, October 2026, with a 1-year Compute Savings Plan for baseline nodes",
                "GitHub Enterprise Cloud licences are paid by Tailspin and not included",
            ],
            # calculator_url from the cost estimate agent. The SOW refers to this link
            # instead of restating the run-rate (see pricing-refers-to-cost-estimate).
            "source_link": "https://calculator.aws/#/estimate?id=sample-tailspin-platform",
            "source_file": None,
        },
    }
    validate(upstream, out / "upstream.json")
    (out / "upstream.json").write_text(json.dumps(upstream, indent=2) + "\n", encoding="utf-8")

    write_text(out / "kickoff-call-transcript.txt", """
Kick-off call transcript - Tailspin Toys DevOps platform
Date: 22 October 2026, 45 minutes, Microsoft Teams
Ingram Micro: Priya Nair (seller), Tomas Novak (DevOps architect)
Tailspin Toys: Daniel Weber (CTO), Sophie Klein (IT Operations Lead), Jonas Becker (Lead Developer, E-commerce)

Priya Nair: Thanks everyone. The migration team is moving the servers to Frankfurt. Today is about how you build and release the shop.
Daniel Weber: Right. We release once a month, on a Saturday, six hours of downtime. Two or three releases a quarter go wrong and we roll back by hand.
Sophie Klein: Two people know the deployment scripts. If one of them is on holiday we do not release.
Tomas Novak: What does the e-commerce platform look like?
Jonas Becker: Nine .NET services. Six are .NET 8 already, three are still .NET Framework 4.8. A React storefront. SQL Server behind it. Code is in Azure DevOps, builds run on one Jenkins box.
Tomas Novak: The six .NET 8 services and the storefront are good candidates for containers on EKS. The three Framework services would stay on Windows EC2 for now, but we can still give them proper pipelines.
Daniel Weber: Agreed. We bought GitHub Enterprise for next year, so we want to move the code and the pipelines to GitHub Actions.
Tomas Novak: Then we would write reusable workflows: build, unit test, code and dependency scanning, image build, image scan, deploy to dev, test and prod with approvals.
Sophie Klein: And infrastructure? Today every environment is a bit different.
Tomas Novak: Everything in Terraform, one set of modules, three environments. No more drift.
Daniel Weber: Christmas is the big one. Mid-November to Christmas Eve we see five times normal traffic. Last year the site was slow for two evenings.
Tomas Novak: Karpenter will add nodes as the load grows. And before the code freeze we run a load test at five times peak, roughly two and a half thousand requests per second.
Daniel Weber: The code freeze starts on the first of November 2027. The platform must be live and tested before that.
Sophie Klein: We have no central logging. Customers tell us when something is broken.
Tomas Novak: Container Insights and CloudWatch Logs, Managed Prometheus and Grafana for metrics and dashboards, alarms to Teams and PagerDuty. And runbooks for the top ten incidents.
Daniel Weber: How do we stop a bad release?
Tomas Novak: Canary releases with Argo Rollouts. A small share of traffic first; if error rate or latency breaks the SLO, it rolls back automatically.
Sophie Klein: GDPR. Everything stays in Frankfurt, logs included.
Tomas Novak: Yes, eu-central-1 only. Secrets go to Secrets Manager, findings to Security Hub, images scanned by Inspector.
Priya Nair: How will you measure success?
Daniel Weber: Weekly releases at least, without downtime. Change failure under ten percent. When something breaks, fixed within the hour.
Jonas Becker: And lead time. Today a one-line fix takes three weeks to reach production. I want it under a day.
Sophie Klein: My platform team is four people. They must be able to run all of this after you leave.
Tomas Novak: We pair with them during the build and finish with handover workshops and runbooks.
Priya Nair: Timeline and budget?
Daniel Weber: Twelve weeks from the first of March. We have about one hundred and fifty thousand euros for delivery.
Priya Nair: Thanks. We will send the statement of work.
""")

    write_xlsx(out / "application-inventory.xlsx", {
        "Applications": [
            ["Service", "Technology", "Repository", "Build", "Deploy method", "Release frequency", "Target", "Owner"],
            ["storefront-web", "React 18", "AzDO: shop/storefront", "Jenkins", "Manual (IIS copy)", "Monthly", "EKS", "Jonas Becker"],
            ["catalog-api", ".NET 8", "AzDO: shop/catalog", "Jenkins", "PowerShell script", "Monthly", "EKS", "Jonas Becker"],
            ["cart-api", ".NET 8", "AzDO: shop/cart", "Jenkins", "PowerShell script", "Monthly", "EKS", "Jonas Becker"],
            ["checkout-api", ".NET 8", "AzDO: shop/checkout", "Jenkins", "PowerShell script", "Monthly", "EKS", "Jonas Becker"],
            ["pricing-api", ".NET 8", "AzDO: shop/pricing", "Jenkins", "PowerShell script", "Monthly", "EKS", "Lena Vogel"],
            ["search-api", ".NET 8", "AzDO: shop/search", "Jenkins", "PowerShell script", "Monthly", "EKS", "Lena Vogel"],
            ["customer-api", ".NET 8", "AzDO: shop/customer", "Jenkins", "PowerShell script", "Monthly", "EKS", "Lena Vogel"],
            ["order-sync", ".NET Framework 4.8", "AzDO: erp/order-sync", "Jenkins", "Manual (MSI)", "Quarterly", "Windows EC2", "Sophie Klein"],
            ["payment-gateway-adapter", ".NET Framework 4.8", "AzDO: shop/payments", "Jenkins", "Manual (MSI)", "Quarterly", "Windows EC2", "Sophie Klein"],
            ["warehouse-feed", ".NET Framework 4.8", "AzDO: erp/warehouse", "Jenkins", "Manual (MSI)", "Quarterly", "Windows EC2", "Sophie Klein"],
        ],
        "Current tooling": [
            ["Area", "Today", "Target"],
            ["Source control", "Azure DevOps Repos", "GitHub Enterprise Cloud"],
            ["CI", "Jenkins (1 server, no backups)", "GitHub Actions"],
            ["CD", "PowerShell scripts, manual", "GitHub Actions + Argo Rollouts (canary)"],
            ["Infrastructure", "Manual, per environment", "Terraform modules"],
            ["Logging", "Local files on servers", "CloudWatch Logs + Container Insights"],
            ["Metrics / dashboards", "None", "Amazon Managed Prometheus + Managed Grafana"],
            ["Alerting", "E-mail from server monitor", "PagerDuty + Microsoft Teams"],
            ["Security scanning", "None", "Inspector, Dependabot, CodeQL, Security Hub"],
            ["Secrets", "web.config files", "AWS Secrets Manager"],
        ],
        "Peak traffic": [
            ["Period", "Requests per second (avg)", "Requests per second (peak)", "Orders per hour (peak)"],
            ["Normal (Jan-Oct)", 180, 500, 900],
            ["Black Friday week", 520, 1800, 3100],
            ["Christmas (15 Nov - 24 Dec)", 650, 2500, 4200],
        ],
    })

    write_diagram(out / "architecture-diagram.png", "Tailspin Toys - DevOps platform on AWS (eu-central-1)", [
        {"label": "Developers", "x": 30, "y": 100, "w": 290, "h": 520, "color": "#2E75B6", "boxes": [
            ("gh", "GitHub", "Enterprise Cloud", 60, 160), ("gha", "GitHub Actions", "build, test, scan", 60, 330),
            ("tf", "Terraform", "IaC, all environments", 60, 500)]},
        {"label": "AWS eu-central-1", "x": 370, "y": 100, "w": 1050, "h": 700, "color": "#E06C00", "boxes": [
            ("sm", "Secrets Manager", "app secrets", 395, 160), ("argo", "Argo Rollouts", "canary + rollback", 655, 160),
            ("ec2", "Windows EC2", ".NET 4.8 services", 915, 160),
            ("ecr", "Amazon ECR", "images + scan", 395, 330), ("eks", "Amazon EKS", "Karpenter, 3 AZs", 655, 330),
            ("alb", "Load Balancer", "ALB ingress", 915, 330),
            ("insp", "Inspector", "vulnerabilities", 395, 500), ("cw", "CloudWatch", "Container Insights", 655, 500),
            ("amp", "Managed Prometheus", "metrics", 915, 500), ("amg", "Managed Grafana", "dashboards", 1175, 500),
            ("sh", "Security Hub", "findings", 395, 670)]},
        {"label": "Users and on-call", "x": 1470, "y": 100, "w": 300, "h": 700, "color": "#548235", "boxes": [
            ("shop", "Shoppers", "5x Christmas peak", 1505, 330), ("pd", "PagerDuty / Teams", "alerts", 1505, 580)]},
    ], [("gh", "gha", ""), ("gha", "ecr", "push"), ("ecr", "eks", ""), ("argo", "eks", ""), ("eks", "sm", ""),
        ("eks", "alb", ""), ("shop", "alb", ""), ("ecr", "insp", ""), ("insp", "sh", ""), ("eks", "cw", ""),
        ("eks", "amp", ""), ("amp", "amg", ""), ("amg", "pd", "alerts")])


# ====================================================================== 03: minimal quick test
def scenario_minimal(out: Path) -> None:
    write_text(out / "project-details.txt", """
Northwind Traders - invoice processing with generative AI (quick test: project details only)

Northwind Traders is a food distributor in Seattle with about 400 employees. The accounts payable team (6 people)
processes about 4,000 supplier invoices per month as PDF e-mail attachments, typing them into Microsoft Dynamics 365
Business Central. It takes about 6 minutes per invoice and 4% have errors.

They want an AI assistant on AWS that reads the invoices (Amazon Textract and Amazon Bedrock), extracts supplier,
invoice number, dates, line items, tax and totals, matches them to purchase orders, and creates a draft purchase
invoice in Business Central for an AP clerk to approve. Exceptions (no PO match, price differences over 2%) go to a
review queue.

Targets: under 1 minute of clerk time per invoice, error rate under 1%. Region us-west-2. Timeline 8 weeks.
Budget about USD 90,000 one-time; expected AWS run-rate about USD 1,200 per month.
""")


SCENARIOS = {
    "01-aiml-fabrikam-insurance": scenario_aiml,
    "02-devops-tailspin-toys": scenario_devops,
    "03-quick-test-northwind-traders": scenario_minimal,
}

# sample.json: what the web UI's "Samples" card and run_sample.py read (form values + a one-line summary).
SAMPLE_META = {
    "01-aiml-fabrikam-insurance": {
        "title": "Fabrikam Insurance: claims FNOL assistant", "customer_name": "Fabrikam Insurance",
        "opportunity_id": "OPP-2026-0517", "template_id": "aiml/ai-ml",
        "summary": "Generative AI claims intake (Bedrock, RAG, Textract). Transcript, notes, diagram and full upstream data.",
    },
    "02-devops-tailspin-toys": {
        "title": "Tailspin Toys: DevOps platform", "customer_name": "Tailspin Toys",
        "opportunity_id": "OPP-2026-0498", "template_id": "devops/devops-standard",
        "summary": "CI/CD, EKS and observability for an e-commerce site. Transcript, app inventory, diagram and upstream data.",
    },
    "03-quick-test-northwind-traders": {
        "title": "Northwind Traders: quick test", "customer_name": "Northwind Traders",
        "opportunity_id": None, "template_id": "aiml/ai-ml",
        "summary": "Invoice processing with generative AI. Project details only: the fastest run.",
    },
}


def main() -> None:
    for name, build in SCENARIOS.items():
        out = HERE / name
        out.mkdir(exist_ok=True)
        build(out)
        (out / "sample.json").write_text(json.dumps(SAMPLE_META[name], indent=2) + "\n", encoding="utf-8")
        files = sorted(p.name for p in out.iterdir() if p.is_file())
        print(f"{name}: {', '.join(files)}")


if __name__ == "__main__":
    main()
